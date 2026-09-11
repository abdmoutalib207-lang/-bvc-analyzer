#!/usr/bin/env python3
"""Conserver un fichier fournisseur ET ce qu'on sait de lui.

CE QUI MANQUE AU DÉPÔT AUJOURD'HUI
──────────────────────────────────
Aucune charge utile de source n'est conservée. Ni réponse HTTP, ni PDF, ni
export. Les séries ont été écrites par le moteur, à des dates et par des
versions de code que la donnée n'enregistre pas. C'est la lacune qui empêche,
aujourd'hui, de répondre à « comment ce prix est-il arrivé là ? ».

Ce module ouvre le premier emplacement où cette réponse pourra exister.

⚠️ L'ANALYSEUR N'EST PAS FIGÉ, ET C'EST VOLONTAIRE
──────────────────────────────────────────────────
    « Avant de figer le format de l'importateur, examine un véritable fichier
      fournisseur. »  — revue externe

Deviner un schéma avant d'avoir vu un fichier produit un analyseur qui décrit
ce qu'on imaginait, pas ce qu'on reçoit. Ce module sépare donc deux choses :

  L'ENVELOPPE — provenance, date de récupération, empreinte, unités déclarées,
                période couverte. Elle ne dépend d'aucun format et elle est
                écrite, testée, utilisable dès maintenant.

  L'ANALYSEUR — la lecture des colonnes. Il est **vide**. `ANALYSEURS` ne
                contient aucune entrée, et `--inspecter` sert précisément à
                regarder un fichier réel avant d'en écrire une.

USAGE
    python pipeline/import_source.py --inspecter mon_export.xlsx
    python pipeline/import_source.py --enregistrer mon_export.xlsx \\
        --titre ADH --fournisseur "Bourse de Casablanca" \\
        --url "https://…" --recupere-le 2026-09-12
"""

from __future__ import annotations

import argparse
import hashlib
import json
import shutil
from datetime import datetime
from pathlib import Path

RACINE = Path(__file__).resolve().parent.parent
SOURCES = RACINE / "sources"

# ⚠️ VIDE À DESSEIN. Une entrée ici signifie « nous savons lire ce format
# parce que nous en avons vu un ». Tant que le dossier `sources/` ne contient
# aucun export réel, il n'y a rien à déclarer — et un analyseur écrit à
# l'aveugle serait une supposition déguisée en code.
ANALYSEURS: dict = {}


def empreinte(chemin: Path) -> str:
    h = hashlib.sha256()
    with chemin.open("rb") as f:
        for bloc in iter(lambda: f.read(1 << 16), b""):
            h.update(bloc)
    return h.hexdigest()


def inspecter(chemin: Path) -> dict:
    """Décrire un fichier SANS supposer son schéma.

    Rend ce qui est observable : taille, empreinte, type apparent, et — si le
    fichier est textuel ou tabulaire — ses premières lignes telles quelles.
    Aucune colonne n'est nommée, aucune unité n'est devinée.
    """
    rapport = {
        "fichier": chemin.name,
        "taille_octets": chemin.stat().st_size,
        "empreinte_sha256": empreinte(chemin),
        "extension": chemin.suffix.lower(),
        "inspecte_le": datetime.now().strftime("%Y-%m-%d %H:%M"),
        "premiers_octets_hex": chemin.open("rb").read(16).hex(),
        "apercu": None,
        "colonnes_apparentes": None,
        "_reserve": "⚠️ Aucun schéma n'est supposé. Les en-têtes éventuels sont "
                    "rapportés TELS QUELS : ce sont des chaînes observées, pas "
                    "des champs reconnus.",
    }

    if chemin.suffix.lower() in (".csv", ".txt", ".tsv"):
        try:
            lignes = chemin.read_text(encoding="utf-8",
                                      errors="replace").splitlines()[:10]
            rapport["apercu"] = lignes
            if lignes:
                for sep in (";", ",", "\t", "|"):
                    if sep in lignes[0]:
                        rapport["colonnes_apparentes"] = lignes[0].split(sep)
                        rapport["separateur_apparent"] = sep
                        break
        except OSError as e:
            rapport["apercu"] = f"lecture impossible : {e}"

    elif chemin.suffix.lower() in (".xlsx", ".xlsm"):
        try:
            import openpyxl
            cl = openpyxl.load_workbook(chemin, read_only=True, data_only=True)
            f = cl[cl.sheetnames[0]]
            lignes = []
            for i, row in enumerate(f.iter_rows(values_only=True)):
                if i >= 10:
                    break
                lignes.append([str(v) if v is not None else None for v in row])
            rapport["feuilles"] = cl.sheetnames
            rapport["apercu"] = lignes
            if lignes:
                rapport["colonnes_apparentes"] = lignes[0]
            cl.close()
        except ImportError:
            rapport["apercu"] = ("openpyxl absent — ajouter la dépendance à "
                                 "requirements_pipeline.txt avant d'importer "
                                 "un XLSX")
        except Exception as e:                       # fichier illisible
            rapport["apercu"] = f"lecture impossible : {type(e).__name__} — {e}"

    return rapport


def enregistrer(chemin: Path, titre: str, fournisseur: str, url: str,
                recupere_le: str, unites: dict | None = None,
                periode: list | None = None, notes: str = "") -> dict:
    """Conserver le fichier ORIGINAL et son enveloppe de provenance.

    Le fichier est copié tel quel, jamais transformé. L'enveloppe l'accompagne
    dans un fichier voisin, pour qu'on ne puisse pas avoir l'un sans l'autre.
    """
    dossier = SOURCES / titre
    dossier.mkdir(parents=True, exist_ok=True)

    # ⚠️ LE NOM D'ARRIVÉE PORTE L'EMPREINTE DU CONTENU.
    # Une version antérieure copiait vers un nom fixe : deux exports différents
    # téléchargés le même jour, portant le même nom de fichier, s'écrasaient
    # l'un l'autre en silence — et l'enveloppe du second décrivait alors un
    # contenu que le premier n'avait plus. Relevé par la revue.
    # Deux fichiers identiques partagent leur empreinte : les réenregistrer ne
    # crée pas de doublon, c'est le comportement voulu.
    emp = empreinte(chemin)
    cible = dossier / f"{chemin.stem}.{emp[:12]}{chemin.suffix}"
    if cible.exists() and empreinte(cible) != emp:
        raise RuntimeError(
            f"collision impossible sur {cible} — même empreinte tronquée, "
            f"contenu différent. Allonger le préfixe d'empreinte.")
    if not cible.exists():
        shutil.copy2(chemin, cible)

    enveloppe = {
        "_quoi": "Fichier fournisseur conservé TEL QUEL, avec sa provenance.",
        "_ce_qui_est_etabli": "l'origine du fichier et son intégrité",
        "_ce_qui_ne_l_est_pas": (
            "⚠️ Conserver un fichier n'établit NI l'exactitude de son contenu, "
            "NI la manière dont le fournisseur a traité les opérations sur "
            "titres. Ce sont des questions distinctes."),
        "titre": titre,
        "fichier": cible.name,
        "fournisseur": fournisseur,
        "url": url,
        "recupere_le": recupere_le,
        "enregistre_le": datetime.now().strftime("%Y-%m-%d %H:%M"),
        "empreinte_sha256": emp,
        "nom_d_origine": chemin.name,
        "taille_octets": cible.stat().st_size,
        "unites_declarees": unites or {
            "prix": "non déclarée",
            "volume": "non déclarée",
            "_note": "⚠️ « non déclarée » est le défaut. Une unité n'est "
                     "inscrite que si le fournisseur la documente — jamais "
                     "déduite de l'ordre de grandeur des nombres.",
        },
        "periode_couverte": periode or [None, None],
        "periode_etablie": bool(periode),
        "analyseur": None,
        "_analyseur_note": (
            "⚠️ Aucun analyseur n'est déclaré. Le format sera figé APRÈS examen "
            "d'un fichier réel, pas avant. Voir --inspecter."),
        "notes": notes,
    }
    (dossier / f"{cible.stem}.provenance.json").write_text(
        json.dumps(enveloppe, ensure_ascii=False, indent=2), encoding="utf-8")
    return enveloppe


def inventaire() -> dict:
    """Ce que `sources/` contient réellement aujourd'hui."""
    fichiers = []
    if SOURCES.exists():
        for p in sorted(SOURCES.rglob("*.provenance.json")):
            fichiers.append(json.loads(p.read_text(encoding="utf-8")))
    return {
        "dossier": str(SOURCES.relative_to(RACINE)),
        "fichiers_conserves": len(fichiers),
        "titres": sorted({f["titre"] for f in fichiers}),
        "analyseurs_declares": sorted(ANALYSEURS),
        "_etat": ("AUCUNE charge utile de fournisseur conservée à ce jour."
                  if not fichiers else
                  f"{len(fichiers)} fichier(s) conservé(s)."),
        "fichiers": fichiers,
    }


def main() -> None:
    ap = argparse.ArgumentParser(description=__doc__)
    ap.add_argument("--inspecter", type=Path)
    ap.add_argument("--enregistrer", type=Path)
    ap.add_argument("--titre")
    ap.add_argument("--fournisseur", default="")
    ap.add_argument("--url", default="")
    ap.add_argument("--recupere-le", default="")
    ap.add_argument("--notes", default="")
    ap.add_argument("--inventaire", action="store_true")
    a = ap.parse_args()

    if a.inspecter:
        print(json.dumps(inspecter(a.inspecter), ensure_ascii=False, indent=2))
        return
    if a.enregistrer:
        if not a.titre:
            ap.error("--titre est requis pour enregistrer")
        e = enregistrer(a.enregistrer, a.titre, a.fournisseur, a.url,
                        a.recupere_le or datetime.now().strftime("%Y-%m-%d"),
                        notes=a.notes)
        print(json.dumps(e, ensure_ascii=False, indent=2))
        return
    print(json.dumps(inventaire(), ensure_ascii=False, indent=2))


if __name__ == "__main__":
    main()
